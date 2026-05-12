from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_MD = ROOT / "docs" / "final-technical-report.md"
REPORT_DOCX = ROOT / "docs" / "final-technical-report.docx"


def add_code_block(document, code):
    for line in code.strip().splitlines():
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8)


def add_markdown_image(document, line):
    alt_start = line.find("[") + 1
    alt_end = line.find("]")
    path_start = line.find("(") + 1
    path_end = line.find(")")
    alt = line[alt_start:alt_end]
    image_path = ROOT / line[path_start:path_end]

    if image_path.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Inches(5.75))

        caption = document.add_paragraph(alt)
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in caption.runs:
            run.italic = True
            run.font.size = Pt(8.5)


def build():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(8.9)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.color.rgb = RGBColor(22, 50, 79)
    styles["Heading 2"].font.name = "Aptos"
    styles["Heading 2"].font.size = Pt(11.5)
    styles["Heading 2"].font.color.rgb = RGBColor(22, 50, 79)

    lines = REPORT_MD.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines = []

    for line in lines:
        if line.startswith("```"):
            if in_code:
                add_code_block(document, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if line.startswith("# "):
            paragraph = document.add_heading(line[2:], level=1)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.add_run(line[2:])
        elif line.startswith("!["):
            add_markdown_image(document, line)
        elif line.strip():
            paragraph = document.add_paragraph(line)
            paragraph.paragraph_format.space_after = Pt(4)

    document.save(REPORT_DOCX)


if __name__ == "__main__":
    build()
